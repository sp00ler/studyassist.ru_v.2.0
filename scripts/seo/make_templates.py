#!/usr/bin/env python3
"""Generate the downloadable Word templates for the /gid guides.

Usage (from the repo root):  python scripts/seo/make_templates.py
Outputs:
  public/templates/shablon-otcheta-po-praktike.docx
  public/templates/dnevnik-praktiki.docx
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "public" / "templates"
FONT = "Times New Roman"


# ---------- helpers ----------

def base_document(size_pt=14, line=1.5):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(30), Mm(15)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(size_pt)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    pf = st.paragraph_format
    pf.line_spacing = line
    pf.space_before = pf.space_after = Pt(0)
    return doc


def para(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, bold=False,
         italic=False, size=None, line=None, after=None, before=None):
    p = doc.add_paragraph()
    p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    if line is not None:
        p.paragraph_format.line_spacing = line
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    return p


def heading(doc, text, page_break=True):
    """Structural heading: caps, centred, bold, no indent. Real Heading 1 style so
    the outline / navigation pane works, restyled to match the body font."""
    p = doc.add_paragraph(style="Heading 1")
    if page_break:
        p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(14)
    r.font.color.rgb = None
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    return p


def add_field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def set_cell_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def cell_text(cell, text, bold=False, italic=False, size=12, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)


# ---------- report template ----------

def build_report():
    doc = base_document()
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True  # title page: no number

    # page number bottom centre (primary footer; first-page footer stays empty)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Cm(0)
    run = fp.add_run()
    run.font.size = Pt(12)
    run.font.name = FONT
    add_field(run, "PAGE")

    C = WD_ALIGN_PARAGRAPH.CENTER
    # --- title page ---
    para(doc, "[МИНИСТЕРСТВО (ВЕДОМСТВО) — ПО ФОРМЕ ВАШЕЙ КАФЕДРЫ]", C, False, size=12, line=1.0, after=6)
    para(doc, "[Название вуза]", C, False, bold=True, line=1.0, after=6)
    para(doc, "[Кафедра]", C, False, line=1.0, after=110)
    para(doc, "ОТЧЁТ", C, False, bold=True, size=16, line=1.0, after=6)
    para(doc, "о прохождении [учебной / производственной / преддипломной] практики", C, False, line=1.15, after=6)
    para(doc, "[наименование организации — места практики]", C, False, line=1.15, after=6)
    para(doc, "[даты начала и окончания практики]", C, False, line=1.15, after=90)

    for label, value in (
        ("Выполнил(а): студент(ка)", "[ФИО студента]"),
        ("Группа:", "[группа]"),
        ("Руководитель практики от вуза:", "[должность, ФИО]"),
        ("Руководитель практики от организации:", "[должность, ФИО]"),
    ):
        p = para(doc, "", WD_ALIGN_PARAGRAPH.LEFT, False, line=1.15, after=6)
        p.paragraph_format.left_indent = Cm(7.5)
        p.add_run(f"{label} ").bold = False
        p.add_run(value)

    p = para(doc, "", C, False, line=1.0, before=70)
    p.add_run("[Город] [год]")

    # --- body ---
    hint = lambda t: para(doc, t, italic=True)

    heading(doc, "СОДЕРЖАНИЕ")
    hint("[Автособираемое содержание: после того как заполните разделы, вставьте его через "
         "«Ссылки» → «Оглавление». Заголовки в шаблоне уже размечены стилем «Заголовок 1».]")
    hint("[Если методичка требует, чтобы в содержании были номера страниц с точками-заполнителями, "
         "Word сделает это сам. Проверьте номера страниц перед сдачей: они сдвигаются при правках.]")

    heading(doc, "ВВЕДЕНИЕ")
    hint("[Актуальность: одним-двумя предложениями, зачем практика нужна на вашем этапе обучения.]")
    hint("[Цель практики и 3–5 задач — берите из программы практики или задания кафедры. "
         "Задачи начинайте с глаголов: изучить, проанализировать, выполнить, подготовить.]")
    hint("[Место и сроки: название организации, подразделение, даты, руководитель от организации.]")

    heading(doc, "1 ХАРАКТЕРИСТИКА ОРГАНИЗАЦИИ")
    hint("[1.1 Общие сведения: полное и краткое название, организационно-правовая форма, год основания, "
         "адрес, вид деятельности. Пишите по открытым данным: устав, сайт организации, реестры.]")
    hint("[1.2 Организационная структура: подразделения и подчинённость. Схему можно вынести в приложение.]")
    hint("[1.3 Основные направления деятельности и показатели за 2–3 года — только те, что вам разрешили использовать.]")

    heading(doc, "2 ВЫПОЛНЕННАЯ РАБОТА")
    hint("[Опишите, что вы делали лично: какие задачи поручали, какими документами, программами и методами "
         "пользовались. Идите по задачам из введения — так текст получается логичным.]")
    hint("[Добавьте анализ или расчёты по заданию кафедры: таблицы, схемы, примеры документов (в приложение). "
         "Каждую таблицу и рисунок подпишите и упомяните в тексте.]")
    hint("[Для преддипломной практики: здесь же — сбор и анализ материала для ВКР.]")

    heading(doc, "ЗАКЛЮЧЕНИЕ")
    hint("[Выводы по каждой задаче из введения: что сделано, к какому результату пришли.]")
    hint("[Какие профессиональные навыки и компетенции получили. Предложения по улучшению работы организации — "
         "если они есть и обоснованы.]")

    heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    hint("[Оформляйте по ГОСТ Р 7.0.100-2018, нумерация сквозная. Внутренние документы организации "
         "(положения, регламенты, инструкции) обычно указывают в списке или в приложении — уточните у руководителя.]")
    para(doc, "1. [Нормативные правовые акты — Конституция, кодексы, федеральные законы.]", indent=False)
    para(doc, "2. [Учебная и научная литература.]", indent=False)
    para(doc, "3. [Электронные ресурсы: название, URL, дата обращения.]", indent=False)

    heading(doc, "ПРИЛОЖЕНИЯ")
    hint("[Приложение А — документы и формы организации (без конфиденциальных данных), приложение Б — таблицы и расчёты, "
         "приложение В — схемы. Каждое приложение с новой страницы, с заголовком «ПРИЛОЖЕНИЕ А» и названием.]")
    hint("[Дневник практики и характеристику от организации обычно сдают отдельно или вкладывают в отчёт — "
         "смотрите методичку кафедры.]")

    return doc


# ---------- diary template ----------

def build_diary():
    doc = base_document(size_pt=12, line=1.15)
    C = WD_ALIGN_PARAGRAPH.CENTER
    L = WD_ALIGN_PARAGRAPH.LEFT

    para(doc, "[Название вуза]", C, False, bold=True, size=14, after=2)
    para(doc, "ДНЕВНИК ПРАКТИКИ", C, False, bold=True, size=16, after=10)

    for label in (
        "Студент(ка): [ФИО]",
        "Группа: [группа]",
        "Направление подготовки / специальность: [код и название]",
        "Вид практики: [учебная / производственная / преддипломная]",
        "Место прохождения: [наименование организации, подразделение]",
        "Сроки практики: с [дд.мм.гггг] по [дд.мм.гггг]",
        "Руководитель от вуза: [должность, ФИО]",
        "Руководитель от организации: [должность, ФИО]",
    ):
        para(doc, label, L, False, after=3)
    para(doc, "", L, False, after=4)

    rows_data = [
        ("[дд.мм.гггг]", "[Пример] Инструктаж по охране труда и технике безопасности. Знакомство с "
         "руководителем, структурой подразделения и режимом работы. Получение индивидуального задания.", "[пример]"),
        ("[дд.мм.гггг]", "[Пример] Изучение внутренних документов: положения о подразделении, должностных "
         "инструкций, регламента документооборота. Получение доступа к рабочим программам.", "[пример]"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_cell_borders(table)
    widths = (Cm(2.8), Cm(9.7), Cm(3.5))

    head = ("Дата", "Содержание выполненной работы", "Подпись руководителя")
    for i, t in enumerate(head):
        cell_text(table.rows[0].cells[i], t, bold=True, align=C)
    # repeat header row on each page
    trPr = table.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)

    for d, c, s in rows_data:
        cells = table.add_row().cells
        cell_text(cells[0], d, italic=True, align=C)
        cell_text(cells[1], c, italic=True)
        cell_text(cells[2], s, italic=True, align=C)
    for _ in range(15):
        row = table.add_row()
        row.height = Cm(1.6)
        for c in row.cells:
            cell_text(c, "")
    for row in table.rows:
        # never split a row across pages
        cant = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant)
        for w, c in zip(widths, row.cells):
            c.width = w

    para(doc, "", L, False, after=6)
    para(doc, "Две первые строки заполнены как пример — удалите их или замените своими записями.",
         L, False, italic=True, size=10, after=14)

    para(doc, "Руководитель практики от организации: ____________________ / [ФИО] /", L, False, after=8)
    para(doc, "Дата: «___» ____________ 20___ г.", L, False, after=8)
    para(doc, "М.П. (если требуется печать организации)", L, False, size=11)
    return doc


# ---------- validation ----------

def validate_report(path):
    d = Document(path)
    texts = [p.text for p in d.paragraphs]
    for h in ("СОДЕРЖАНИЕ", "ВВЕДЕНИЕ", "1 ХАРАКТЕРИСТИКА ОРГАНИЗАЦИИ", "2 ВЫПОЛНЕННАЯ РАБОТА",
              "ЗАКЛЮЧЕНИЕ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "ПРИЛОЖЕНИЯ"):
        assert h in texts, f"heading missing: {h}"
    s = d.sections[0]
    assert s.different_first_page_header_footer
    assert round(s.left_margin.mm) == 30 and round(s.right_margin.mm) == 15
    assert round(s.top_margin.mm) == 20 and round(s.bottom_margin.mm) == 20
    assert "PAGE" in s.footer._element.xml
    assert "PAGE" not in s.first_page_footer._element.xml
    print(f"  OK {path.name}: {len(d.paragraphs)} paragraphs, margins 30/15/20/20, PAGE field in footer")


def validate_diary(path):
    d = Document(path)
    assert len(d.tables) == 1
    t = d.tables[0]
    assert len(t.columns) == 3, "diary table must have 3 columns"
    assert [c.text for c in t.rows[0].cells] == ["Дата", "Содержание выполненной работы", "Подпись руководителя"]
    assert len(t.rows) >= 17
    assert any("ДНЕВНИК ПРАКТИКИ" in p.text for p in d.paragraphs)
    print(f"  OK {path.name}: table {len(t.rows)} rows x {len(t.columns)} columns")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    r = OUT / "shablon-otcheta-po-praktike.docx"
    dn = OUT / "dnevnik-praktiki.docx"
    build_report().save(r)
    build_diary().save(dn)
    print("Generated:")
    validate_report(r)
    validate_diary(dn)
    for f in (r, dn):
        print(f"  {f.relative_to(ROOT)}  {f.stat().st_size} bytes")


if __name__ == "__main__":
    main()
